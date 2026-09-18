#!/usr/bin/env python3
"""客户交付 Excel 工具：统一样式、说明页、列名映射、值翻译。

作为模块导入使用：

    from xlsx_kit import Sheet, ReadmeSheet, write_workbook, tsv_to_sheet

    readme = ReadmeSheet("说明", title="……", intro=[...], sheets_desc=[("总表", "…")], columns_desc={...})
    sheet = tsv_to_sheet("总表", "work/final/x.tsv", columns=[
        Column("query_gene", "品种 A 基因 ID", width=22),
        Column("tier", "置信度", mapping={"HC": "高", "MC": "中", "LC": "低"}, width=10),
    ], row_filter=lambda r: r["tier"] in ("HC", "MC"))
    write_workbook("交付.xlsx", [readme, sheet])

约定：
- 每个工作簿第一页是「说明」页：项目、内容、每张表和每列的含义、使用建议。
- 列名一律用客户能懂的中文；内部列名只出现在 Column.src，绝不写进单元格。
- 缺失值写空单元格，不写 NA / None / 0。
- 首行冻结 + 自动筛选；表头青绿底白字；列宽按内容估算并封顶。
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="087F73")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="075D55")
SUBTITLE_FONT = Font(bold=True, size=11, color="203B38")
MUTED_FONT = Font(color="5F726F")
THIN = Side(style="thin", color="DBE5E2")
BORDER = Border(bottom=THIN)

EXCEL_MAX_ROWS = 1_048_576


@dataclass
class Column:
    src: str                       # 源列名（内部），不会出现在输出里
    header: str                    # 输出列名（客户可读）
    mapping: dict | None = None    # 值翻译，如 {"HC": "高"}
    fmt: Callable | None = None    # 值格式化函数
    width: int | None = None       # 列宽（字符），None 自动
    number_format: str | None = None  # openpyxl 数字格式，如 "0.0%"
    desc: str = ""                 # 列含义（写入说明页）


@dataclass
class Sheet:
    name: str
    headers: Sequence[str]
    rows: Iterable[Sequence]
    widths: Sequence[int | None] | None = None
    number_formats: Sequence[str | None] | None = None
    freeze: bool = True
    autofilter: bool = True
    desc: str = ""                                  # 表含义（写入说明页）
    column_desc: Sequence[tuple[str, str]] = field(default_factory=list)  # (列名, 含义)


@dataclass
class ReadmeSheet:
    name: str = "说明"
    title: str = ""
    intro: Sequence[str] = field(default_factory=list)          # 段落
    sheets_desc: Sequence[tuple[str, str]] = field(default_factory=list)   # (表名, 内容)
    columns_desc: dict[str, Sequence[tuple[str, str]]] = field(default_factory=dict)  # 表名 -> [(列, 含义)]
    notes: Sequence[str] = field(default_factory=list)          # 使用建议 / 注意
    footer: str = ""


def _safe_sheet_name(name: str) -> str:
    bad = '[]:*?/\\'
    for ch in bad:
        name = name.replace(ch, " ")
    return name[:31]


def tsv_to_sheet(
    name: str,
    path: str | Path,
    columns: Sequence[Column],
    row_filter: Callable[[dict], bool] | None = None,
    sort_key: Callable[[dict], object] | None = None,
    desc: str = "",
    delimiter: str = "\t",
) -> Sheet:
    """按列规范读取 TSV，输出客户可读的 Sheet。"""
    with open(path, encoding="utf-8", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=delimiter)
        data = [r for r in reader if row_filter is None or row_filter(r)]
    if sort_key:
        data.sort(key=sort_key)
    rows = []
    for r in data:
        out = []
        for c in columns:
            v = r.get(c.src, "")
            if c.mapping is not None:
                v = c.mapping.get(v, v)
            if c.fmt is not None:
                v = c.fmt(v)
            out.append(v)
        rows.append(out)
    return Sheet(
        name=name,
        headers=[c.header for c in columns],
        rows=rows,
        widths=[c.width for c in columns],
        number_formats=[c.number_format for c in columns],
        desc=desc,
        column_desc=[(c.header, c.desc) for c in columns if c.desc],
    )


def _estimate_width(values: list, header: str, cap: int = 48) -> int:
    def w(s):
        s = "" if s is None else str(s)
        return sum(2 if ord(ch) > 127 else 1 for ch in s)

    sample = values[:500]
    longest = max([w(header)] + [w(v) for v in sample]) if sample else w(header)
    return max(8, min(cap, longest + 2))


def _write_sheet(ws, sheet: Sheet) -> int:
    ws.append(list(sheet.headers))
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    n = 0
    rows = list(sheet.rows)
    if len(rows) + 1 > EXCEL_MAX_ROWS:
        raise ValueError(f"工作表 {sheet.name} 行数 {len(rows)} 超过 Excel 上限，请拆分或改为 CSV 交付")
    for r in rows:
        ws.append(["" if v is None else v for v in r])
        n += 1
    ncol = len(sheet.headers)
    for i in range(ncol):
        col_values = [r[i] for r in rows] if rows else []
        width = (sheet.widths[i] if sheet.widths and sheet.widths[i] else None) or _estimate_width(col_values, sheet.headers[i])
        ws.column_dimensions[get_column_letter(i + 1)].width = width
        fmt = sheet.number_formats[i] if sheet.number_formats else None
        if fmt:
            for cell in ws.iter_cols(min_col=i + 1, max_col=i + 1, min_row=2, max_row=n + 1):
                for c in cell:
                    c.number_format = fmt
    if sheet.freeze:
        ws.freeze_panes = "A2"
    if sheet.autofilter and n:
        ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{n + 1}"
    return n


def _write_readme(ws, rd: ReadmeSheet, sheets: Sequence[Sheet]):
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    row = 1
    if rd.title:
        ws.cell(row=row, column=1, value=rd.title).font = TITLE_FONT
        row += 2
    for para in rd.intro:
        c = ws.cell(row=row, column=1, value=para)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.row_dimensions[row].height = max(18, 16 * (1 + len(para) // 60))
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="工作表说明").font = SUBTITLE_FONT
    row += 1
    desc_map = dict(rd.sheets_desc)
    for s in sheets:
        ws.cell(row=row, column=1, value=s.name).font = Font(bold=True)
        c = ws.cell(row=row, column=2, value=desc_map.get(s.name, s.desc))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    row += 1

    col_desc = dict(rd.columns_desc)
    for s in sheets:
        cd = col_desc.get(s.name) or s.column_desc
        if not cd:
            continue
        ws.cell(row=row, column=1, value=f"「{s.name}」列说明").font = SUBTITLE_FONT
        row += 1
        for col, meaning in cd:
            ws.cell(row=row, column=1, value=col).font = Font(bold=True)
            c = ws.cell(row=row, column=2, value=meaning)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        row += 1

    if rd.notes:
        ws.cell(row=row, column=1, value="使用建议").font = SUBTITLE_FONT
        row += 1
        for i, n in enumerate(rd.notes, 1):
            c = ws.cell(row=row, column=1, value=f"{i}. {n}")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            ws.row_dimensions[row].height = max(18, 16 * (1 + len(n) // 60))
            row += 1
        row += 1
    if rd.footer:
        ws.cell(row=row, column=1, value=rd.footer).font = MUTED_FONT


def write_workbook(output: str | Path, sheets: Sequence[ReadmeSheet | Sheet]) -> dict[str, int]:
    """写出工作簿；返回每张数据表的行数。ReadmeSheet 必须放在列表首位（若有）。"""
    wb = Workbook()
    wb.remove(wb.active)
    counts: dict[str, int] = {}
    data_sheets = [s for s in sheets if isinstance(s, Sheet)]
    for s in sheets:
        ws = wb.create_sheet(_safe_sheet_name(s.name))
        if isinstance(s, ReadmeSheet):
            _write_readme(ws, s, data_sheets)
        else:
            counts[s.name] = _write_sheet(ws, s)
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return counts
