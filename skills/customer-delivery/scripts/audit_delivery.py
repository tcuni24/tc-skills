#!/usr/bin/env python3
"""交付目录体检：只交付该交付的文件，且文件内容不泄漏内部信息。

用法：
  audit_delivery.py <交付目录> [--allow-ext xlsx,html,png,pdf] [--allow-dir plot]
                    [--terms 词1,词2] [--terms-file 词表.txt] [--allow-terms 词1,词2]

检查项：
  1. 文件清单：只允许白名单扩展名 / 子目录；报告其他一切文件（日志、tsv、脚本、配置、隐藏文件）。
  2. 内容泄漏：扫描 HTML 文本、XLSX 单元格、文件名，命中以下任一即报告：
     - 绝对路径（/data_0/…、/public/…、/home/…、/project/…、/mnt/…、/tmp/…）、Windows 盘符路径
     - 内部工作目录片段（work/、test/、logs/、tmp/）、脚本与日志文件名（*.sh *.py *.log）
     - 当前用户名、主机名
     - 邮箱、IP 地址
     - 开发标记（TODO / FIXME / XXX / 冒烟 / smoke / debug）
     - 自定义术语（--terms：内部列名、内部样本编号、工具内部枚举值等）
  3. HTML 引用：<img src> 必须是相对路径且文件存在；不允许 http(s):// 外部资源。
退出码：有任何问题为 1，否则 0。
"""
from __future__ import annotations

import argparse
import getpass
import html as html_mod
import os
import re
import socket
import sys
import zipfile
from pathlib import Path

DEFAULT_PATTERNS = {
    "绝对路径": re.compile(r"(?<![\w/])/(?:data_\d+|data|public|home|project|mnt|tmp|opt|srv|scratch|work|root|usr/local)/[\w./-]+"),
    "Windows 路径": re.compile(r"\b[A-Za-z]:\\[\w\\.-]+"),
    "内部目录片段": re.compile(r"(?<![\w/])(?:work|test|tests|logs|tmp|scratch|intermediate)/[\w./-]*"),
    "脚本或日志文件名": re.compile(r"\b[\w-]+\.(?:sh|py|log|yml|yaml|smk|nf)\b"),
    "邮箱": re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+"),
    "IP 地址": re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b"),
    "开发标记": re.compile(r"\b(?:TODO|FIXME|XXX|HACK|smoke ?test|debug)\b|冒烟|调试用|内部(?:文件|路径|测试|使用)"),
}


def strip_html(raw: str) -> str:
    raw = re.sub(r"<script.*?</script>|<style.*?</style>", " ", raw, flags=re.S)
    raw = re.sub(r'data:image/[^"\')]+', " ", raw)
    raw = re.sub(r"<[^>]+>", " ", raw)
    return html_mod.unescape(raw)


def xlsx_texts(path: Path):
    """读取 xlsx 中所有字符串单元格（含说明页），以 (sheet, text) 生成。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("警告：未安装 openpyxl，跳过 xlsx 内容扫描", file=sys.stderr)
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and v:
                    yield ws.title, v
    wb.close()


def scan_text(label: str, txt: str, patterns: dict, terms: list[str], allow: set[str], findings: list):
    for name, pat in patterns.items():
        for m in pat.finditer(txt):
            s = m.group(0)
            if s in allow or any(a in s for a in allow):
                continue
            findings.append((label, name, s))
    for t in terms:
        n = txt.count(t)
        if n:
            findings.append((label, "自定义术语", f"{t} ×{n}"))


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("delivery", type=Path)
    ap.add_argument("--allow-ext", default="xlsx,html,png,pdf", help="允许的文件扩展名")
    ap.add_argument("--allow-dir", default="plot", help="允许的子目录（逗号分隔）")
    ap.add_argument("--allow-file", default="", help="额外允许的顶层文件名（如 README.txt,SHA256SUMS）")
    ap.add_argument("--terms", default="", help="逗号分隔的内部术语 / 列名 / 枚举值")
    ap.add_argument("--terms-file", type=Path, default=None, help="每行一个术语")
    ap.add_argument("--allow-terms", default="", help="允许出现的例外字符串（如软件名 diamond.py）")
    ap.add_argument("--max-show", type=int, default=40)
    args = ap.parse_args()

    root = args.delivery.resolve()
    if not root.is_dir():
        print(f"目录不存在: {root}")
        return 2
    allow_ext = {e.strip().lower().lstrip(".") for e in args.allow_ext.split(",") if e.strip()}
    allow_dirs = {d.strip() for d in args.allow_dir.split(",") if d.strip()}
    allow_files = {f.strip() for f in args.allow_file.split(",") if f.strip()}
    terms = [t.strip() for t in args.terms.split(",") if t.strip()]
    if args.terms_file:
        terms += [l.strip() for l in args.terms_file.read_text(encoding="utf-8").splitlines() if l.strip() and not l.startswith("#")]
    allow = {a.strip() for a in args.allow_terms.split(",") if a.strip()}

    patterns = dict(DEFAULT_PATTERNS)
    user = getpass.getuser()
    host = socket.gethostname()
    ident = [x for x in {user, host, host.split(".")[0]} if x and len(x) >= 3]
    if ident:
        patterns["用户名/主机名"] = re.compile("|".join(re.escape(x) for x in ident))

    problems: list[str] = []
    findings: list[tuple] = []
    files: list[Path] = []

    # 1. 文件清单
    for p in sorted(root.rglob("*")):
        rel = p.relative_to(root)
        if p.is_dir():
            if rel.parts[0] not in allow_dirs:
                problems.append(f"非白名单子目录: {rel}/")
            continue
        files.append(p)
        if p.name.startswith("."):
            problems.append(f"隐藏文件: {rel}")
            continue
        if len(rel.parts) > 1 and rel.parts[0] not in allow_dirs:
            problems.append(f"非白名单目录中的文件: {rel}")
            continue
        if p.name in allow_files:
            continue
        if p.suffix.lower().lstrip(".") not in allow_ext:
            problems.append(f"非白名单类型文件: {rel}")
        if p.stat().st_size == 0:
            problems.append(f"空文件: {rel}")
        scan_text(f"文件名 {rel}", str(rel), {k: v for k, v in patterns.items() if k not in ("脚本或日志文件名",)}, terms, allow, findings)

    # 2/3. 内容
    for p in files:
        rel = p.relative_to(root)
        suf = p.suffix.lower()
        if suf in (".html", ".htm"):
            raw = p.read_text(encoding="utf-8", errors="replace")
            for m in re.finditer(r'(?:src|href)=["\'](https?://[^"\']+)', raw):
                problems.append(f"{rel}: 外部资源引用 {m.group(1)}")
            for m in re.finditer(r'<img[^>]+src=["\']([^"\']+)["\']', raw):
                src = m.group(1)
                if src.startswith("data:"):
                    continue
                if src.startswith("/") or re.match(r"^[A-Za-z]:", src):
                    problems.append(f"{rel}: 图片使用绝对路径 {src}")
                elif not (p.parent / src).is_file():
                    problems.append(f"{rel}: 图片文件不存在 {src}")
            scan_text(f"HTML {rel}", strip_html(raw), patterns, terms, allow, findings)
        elif suf == ".xlsx":
            try:
                for sheet, v in xlsx_texts(p):
                    scan_text(f"XLSX {rel}[{sheet}]", v, patterns, terms, allow, findings)
            except zipfile.BadZipFile:
                problems.append(f"{rel}: 不是有效的 xlsx")
        elif suf in (".txt", ".md", ".csv", ".tsv"):
            scan_text(f"文本 {rel}", p.read_text(encoding="utf-8", errors="replace"), patterns, terms, allow, findings)

    # 汇总
    print(f"交付目录: {root}")
    print(f"文件数: {len(files)}")
    for p in files:
        print(f"  {p.relative_to(root)}  ({p.stat().st_size / 1024:.0f} KB)")
    print()
    if problems:
        print(f"文件 / 引用问题 {len(problems)} 项：")
        for x in problems[: args.max_show]:
            print("  - " + x)
    else:
        print("文件 / 引用问题: 无")

    # 聚合泄漏
    agg: dict[tuple, int] = {}
    for label, kind, s in findings:
        agg[(label, kind, s)] = agg.get((label, kind, s), 0) + 1
    if agg:
        print(f"\n内容泄漏疑点 {len(agg)} 项（去重后）：")
        for (label, kind, s), n in list(sorted(agg.items()))[: args.max_show]:
            print(f"  - [{kind}] {label}: {s!r}" + (f" ×{n}" if n > 1 else ""))
        if len(agg) > args.max_show:
            print(f"  … 另有 {len(agg) - args.max_show} 项")
    else:
        print("内容泄漏疑点: 无")

    return 1 if (problems or agg) else 0


if __name__ == "__main__":
    sys.exit(main())
