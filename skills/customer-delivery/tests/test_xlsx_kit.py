#!/usr/bin/env python3
"""Unit tests for xlsx_kit.py export results and boundary behaviors."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from openpyxl import load_workbook  # noqa: E402
from xlsx_kit import (  # noqa: E402
    Column,
    EXCEL_MAX_ROWS,
    ReadmeSheet,
    Sheet,
    tsv_to_sheet,
    write_workbook,
)


class XlsxKitTests(unittest.TestCase):
    """Test suite for xlsx_kit workbook generation and edge case handling."""

    def test_sheet_title_sanitization_and_truncation(self):
        # 工作表名称最长 31 字符，且非法字符 []:*?/\ 必须被替换
        long_and_illegal = "测试[工作表]:名称?超长字符截断测试超过三十一个汉字长度"
        sheet = Sheet(name=long_and_illegal, headers=["列1"], rows=[["值1"]])

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            xlsx_path = Path(f.name)

        try:
            write_workbook(xlsx_path, [sheet])
            wb = load_workbook(xlsx_path, read_only=True)
            actual_name = wb.sheetnames[0]
            self.assertLessEqual(len(actual_name), 31)
            for ch in "[]:*?/\\":
                self.assertNotIn(ch, actual_name)
            wb.close()
        finally:
            xlsx_path.unlink(missing_ok=True)

    def test_tsv_to_sheet_and_empty_value_handling(self):
        # 缺失值写空单元格（None/""），不写 NA / None / 0
        tsv_content = (
            "internal_id\tinternal_score\tinternal_tier\tinternal_note\n"
            "gene_01\t0.95\tHC\t正常\n"
            "gene_02\tNA\tLC\t\n"
            "gene_03\t\tMC\t\n"
        )
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", suffix=".tsv", delete=False) as f:
            f.write(tsv_content)
            tsv_path = Path(f.name)

        try:
            cols = [
                Column("internal_id", "基因ID"),
                Column("internal_score", "分值", fmt=lambda x: float(x) if x and x != "NA" else None),
                Column("internal_tier", "置信度", mapping={"HC": "高", "MC": "中", "LC": "低"}),
                Column("internal_note", "备注", fmt=lambda x: x if x else None),
            ]
            sheet = tsv_to_sheet("基因表", tsv_path, cols)

            self.assertEqual(sheet.headers, ["基因ID", "分值", "置信度", "备注"])
            self.assertEqual(len(sheet.rows), 3)

            # 验证空值/NA 处理
            row2 = sheet.rows[1]
            self.assertEqual(row2[0], "gene_02")
            self.assertIsNone(row2[1])  # "NA" -> None
            self.assertEqual(row2[2], "低")  # LC -> 低
            self.assertIsNone(row2[3])  # empty string -> None

            row3 = sheet.rows[2]
            self.assertIsNone(row3[1])  # empty string -> None
            self.assertIsNone(row3[3])  # empty string -> None
        finally:
            tsv_path.unlink(missing_ok=True)

    def test_tsv_to_sheet_row_filter(self):
        # 仅放过滤后的行
        tsv_content = (
            "id\ttier\n"
            "g1\tHC\n"
            "g2\tLC\n"
            "g3\tMC\n"
        )
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", suffix=".tsv", delete=False) as f:
            f.write(tsv_content)
            tsv_path = Path(f.name)

        try:
            cols = [Column("id", "ID"), Column("tier", "等级")]
            sheet = tsv_to_sheet("过滤表", tsv_path, cols, row_filter=lambda r: r["tier"] in ("HC", "MC"))
            self.assertEqual(len(sheet.rows), 2)
            self.assertEqual(sheet.rows[0], ["g1", "HC"])
            self.assertEqual(sheet.rows[1], ["g3", "MC"])
        finally:
            tsv_path.unlink(missing_ok=True)

    def test_full_workbook_with_readme_and_formatting(self):
        # 验证 ReadmeSheet 与数据表共同写入 Excel
        readme = ReadmeSheet(
            title="示例项目交付说明",
            intro=["本文件为交付结果汇总", "请参考说明页使用"],
            sheets_desc=[("核心表", "建议使用的核心位点")],
            columns_desc={"核心表": [("位点", "位点标识符"), ("质量", "判定等级")]},
        )
        data_sheet = Sheet(
            name="核心表",
            headers=["位点", "质量"],
            rows=[["Site_A", "高质量"], ["Site_B", "中质量"]],
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            xlsx_path = Path(f.name)

        try:
            write_workbook(xlsx_path, [readme, data_sheet])

            wb = load_workbook(xlsx_path, data_only=True)
            self.assertEqual(wb.sheetnames, ["说明", "核心表"])

            # 验证说明页内容
            ws_readme = wb["说明"]
            self.assertEqual(ws_readme["A1"].value, "示例项目交付说明")

            # 验证数据页内容与首行冻结
            ws_data = wb["核心表"]
            self.assertEqual(ws_data["A1"].value, "位点")
            self.assertEqual(ws_data["B1"].value, "质量")
            self.assertEqual(ws_data["A2"].value, "Site_A")
            self.assertEqual(ws_data.freeze_panes, "A2")
            self.assertIsNotNone(ws_data.auto_filter.ref)

            wb.close()
        finally:
            xlsx_path.unlink(missing_ok=True)

    def test_excel_max_rows_constant(self):
        self.assertEqual(EXCEL_MAX_ROWS, 1_048_576)


if __name__ == "__main__":
    unittest.main()
